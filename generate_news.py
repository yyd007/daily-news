#!/usr/bin/env python3
"""Fetch top 10 worldwide, China, and AI headlines and write a dated Word brief."""

from __future__ import annotations

import html
import json
import re
import sys
from datetime import datetime
from email.utils import parsedate_to_datetime
from pathlib import Path
from urllib.parse import parse_qs, urlparse
from zoneinfo import ZoneInfo

import feedparser
import requests
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import html as lxml_html

LOCAL_TZ = ZoneInfo("Asia/Shanghai")
HERE = Path(__file__).resolve().parent
OUTPUT_DIR = HERE
USER_AGENT = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36"
)
HEADERS = {"User-Agent": USER_AGENT, "Accept-Language": "en-US,en;q=0.9,zh-CN;q=0.8"}

NAVY = RGBColor(0x1B, 0x3A, 0x4B)
TEAL = RGBColor(0x1F, 0x6F, 0x6A)
SLATE = RGBColor(0x4A, 0x55, 0x5A)
MUTED = RGBColor(0x6B, 0x72, 0x80)
LINK_BLUE = RGBColor(0x1D, 0x4E, 0x89)

SECTIONS = [
    {
        "key": "worldwide",
        "title": "Top 10 Worldwide",
        "title_zh": "\u5168\u7403\u5934\u6761",
        "feeds": [
            "https://news.google.com/rss/headlines/section/topic/WORLD?hl=en-US&gl=US&ceid=US:en",
            "https://feeds.bbci.co.uk/news/world/rss.xml",
            "https://news.google.com/rss?hl=en-US&gl=US&ceid=US:en",
        ],
    },
    {
        "key": "china",
        "title": "Top 10 in China",
        "title_zh": "\u4e2d\u56fd\u5934\u6761",
        "feeds": [
            "https://news.google.com/rss?hl=zh-CN&gl=CN&ceid=CN:zh-CN",
            "https://news.google.com/rss/search?q=%E4%B8%AD%E5%9B%BD+when:1d&hl=zh-CN&gl=CN&ceid=CN:zh-CN",
            "https://feeds.bbci.co.uk/zhongwen/simp/rss.xml",
            "https://news.google.com/rss/search?q=China+when:1d&hl=en-US&gl=US&ceid=US:en",
        ],
    },
    {
        "key": "ai",
        "title": "Top 10 Related to AI",
        "title_zh": "\u4eba\u5de5\u667a\u80fd",
        "feeds": [
            "https://news.google.com/rss/search?q=artificial+intelligence+OR+%22generative+AI%22+OR+ChatGPT+OR+OpenAI+OR+LLM+when:1d&hl=en-US&gl=US&ceid=US:en",
            "https://news.google.com/rss/search?q=%E4%BA%BA%E5%B7%A5%E6%99%BA%E8%83%BD+OR+%E5%A4%A7%E6%A8%A1%E5%9E%8B+when:1d&hl=zh-CN&gl=CN&ceid=CN:zh-CN",
            "https://www.theverge.com/rss/ai-artificial-intelligence/index.xml",
        ],
    },
]


def fetch_feed(url: str) -> list[dict]:
    try:
        response = requests.get(url, headers=HEADERS, timeout=20)
        response.raise_for_status()
    except requests.RequestException as exc:
        print(f"  skip {url}: {exc}", file=sys.stderr)
        return []

    parsed = feedparser.parse(response.content)
    feed_title = clean_text(parsed.feed.get("title", ""))
    items = []
    for entry in parsed.entries:
        title = clean_text(entry.get("title", ""))
        if not title:
            continue
        link = unwrap_link(entry)
        source = extract_source(entry, title, feed_title, link)
        title = strip_source_suffix(title, source)
        items.append(
            {
                "title": title,
                "link": link,
                "source": source,
                "published": parse_published(entry),
                "summary": clean_summary(entry.get("summary") or entry.get("description") or ""),
            }
        )
    return items


def clean_text(value: str) -> str:
    text = html.unescape(re.sub(r"<[^>]+>", " ", value or ""))
    return re.sub(r"\s+", " ", text).strip()


def clean_summary(value: str) -> str:
    text = clean_text(value)
    text = re.sub(r"^\s*View Full Coverage on Google News\s*", "", text, flags=re.I)
    if len(text) > 280:
        clipped = text[:277].rsplit(" ", 1)[0]
        text = clipped + "..."
    return text


def is_english(text: str) -> bool:
    if not text:
        return False
    cjk = len(re.findall(r"[\u4e00-\u9fff]", text))
    latin = len(re.findall(r"[A-Za-z]", text))
    return latin >= 8 and cjk * 2 < latin


_TRANSLATE_CACHE: dict[str, str] = {}


def _parse_google_translation(payload) -> str:
    if isinstance(payload, str):
        return payload.strip()
    if isinstance(payload, dict):
        return clean_text(str(payload.get("responseData", {}).get("translatedText") or ""))
    if isinstance(payload, list) and payload:
        first = payload[0]
        if isinstance(first, str):
            return first.strip()
        if isinstance(first, list):
            chunks = []
            for part in first:
                if isinstance(part, str) and not re.fullmatch(r"[a-z]{2}(?:-[A-Z]{2})?", part):
                    chunks.append(part)
                elif isinstance(part, list) and part and isinstance(part[0], str):
                    chunks.append(part[0])
            return "".join(chunks).strip()
    return ""


def translate_to_zh(text: str) -> str:
    text = clean_text(text)
    if not is_english(text):
        return ""
    if text in _TRANSLATE_CACHE:
        return _TRANSLATE_CACHE[text]

    endpoints = [
        (
            "https://clients5.google.com/translate_a/t",
            {"client": "dict-chrome-ex", "sl": "auto", "tl": "zh-CN", "q": text},
        ),
        (
            "https://api.mymemory.translated.net/get",
            {"q": text, "langpair": "en|zh-CN"},
        ),
    ]
    translated = ""
    for url, params in endpoints:
        try:
            response = requests.get(url, params=params, headers=HEADERS, timeout=20)
            if response.status_code == 429:
                continue
            response.raise_for_status()
            translated = _parse_google_translation(response.json())
            if translated:
                break
        except (requests.RequestException, json.JSONDecodeError, TypeError, IndexError, KeyError) as exc:
            print(f"  skip translation via {urlparse(url).netloc}: {exc}", file=sys.stderr)

    if translated and translated.lower() == text.lower():
        translated = ""
    _TRANSLATE_CACHE[text] = translated
    return translated


def add_translations(sections: dict[str, list[dict]]) -> None:
    for stories in sections.values():
        for story in stories:
            story["title_zh"] = translate_to_zh(story["title"])
            story["summary_zh"] = ""


DOMAIN_SOURCES = {
    "bbc.co.uk": "BBC News",
    "bbc.com": "BBC News",
    "theverge.com": "The Verge",
    "reuters.com": "Reuters",
    "apnews.com": "AP",
    "nytimes.com": "The New York Times",
    "washingtonpost.com": "The Washington Post",
    "theguardian.com": "The Guardian",
    "cnn.com": "CNN",
    "nbcnews.com": "NBC News",
    "abcnews.go.com": "ABC News",
    "abcnews.com": "ABC News",
    "cbsnews.com": "CBS News",
    "bloomberg.com": "Bloomberg",
    "wsj.com": "The Wall Street Journal",
    "ft.com": "Financial Times",
    "scmp.com": "South China Morning Post",
    "xinhuanet.com": "Xinhua",
    "news.cn": "Xinhua",
    "people.com.cn": "People's Daily",
    "thepaper.cn": "The Paper",
    "chinadaily.com.cn": "China Daily",
    "caixin.com": "Caixin",
    "yicai.com": "Yicai",
    "36kr.com": "36Kr",
    "techcrunch.com": "TechCrunch",
    "wired.com": "Wired",
    "arstechnica.com": "Ars Technica",
    "tomshardware.com": "Tom's Hardware",
}

BAD_SOURCES = {
    "",
    "unknown",
    "unknown source",
    "google news",
    "google ??",
    "???? - google ??",
    "top stories - google news",
    "world - latest - google news",
}


def clean_source_name(value: str) -> str:
    source = clean_text(value)
    source = re.sub(r"\s*-\s*Breaking News, Latest News and Videos$", "", source, flags=re.I)
    source = re.sub(r"\s*-\s*Google ??$", "", source, flags=re.I)
    source = re.sub(r"\s*-\s*Google News$", "", source, flags=re.I)
    source = re.sub(r"\s*\|\s*Latest.*$", "", source, flags=re.I)
    return source.strip(" -|")


def source_from_url(url: str) -> str:
    host = urlparse(url).netloc.lower()
    host = re.sub(r"^(www|m|rss|feeds)\.", "", host)
    for domain, name in DOMAIN_SOURCES.items():
        if host == domain or host.endswith("." + domain):
            return name
    return ""


def extract_source(entry: dict, title: str, feed_title: str, link: str) -> str:
    candidates: list[str] = []
    if hasattr(entry, "source") and getattr(entry.source, "title", None):
        candidates.append(clean_source_name(entry.source.title))
    if " - " in title:
        candidates.append(clean_source_name(title.rsplit(" - ", 1)[-1]))
    if feed_title:
        candidates.append(clean_source_name(feed_title))
    url_source = source_from_url(link)
    if url_source:
        candidates.append(url_source)

    for candidate in candidates:
        if candidate.lower() not in BAD_SOURCES:
            return candidate
    return "Unknown source"


def strip_source_suffix(title: str, source: str) -> str:
    if source and f" - {source}" in title:
        title = title.split(f" - {source}", 1)[0].strip()
    title = re.sub(r"\s*-\s*Breaking News, Latest News and Videos$", "", title, flags=re.I)
    return title


def unwrap_link(entry: dict) -> str:
    link = entry.get("link") or ""
    summary = entry.get("summary") or entry.get("description") or ""
    if summary:
        try:
            tree = lxml_html.fromstring(summary)
            hrefs = [a.get("href") for a in tree.xpath("//a[@href]") if a.get("href")]
            for href in hrefs:
                if "news.google.com" not in href:
                    return href
        except Exception:
            pass
    parsed = urlparse(link)
    if parsed.netloc.endswith("news.google.com"):
        query_url = parse_qs(parsed.query).get("url", [None])[0]
        if query_url:
            return query_url
    return link


def parse_published(entry: dict) -> datetime | None:
    raw = entry.get("published") or entry.get("updated") or ""
    if not raw:
        return None
    try:
        dt = parsedate_to_datetime(raw)
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=LOCAL_TZ)
        return dt.astimezone(LOCAL_TZ)
    except (TypeError, ValueError, OverflowError):
        return None


def normalize_title(title: str) -> str:
    return re.sub(r"[^\w\u4e00-\u9fff]+", "", title.lower())


def collect_top10(feeds: list[str]) -> list[dict]:
    seen: set[str] = set()
    ranked: list[dict] = []
    for url in feeds:
        print(f"  fetching {url}")
        for item in fetch_feed(url):
            key = normalize_title(item["title"])
            if len(key) < 8 or key in seen:
                continue
            seen.add(key)
            ranked.append(item)
            if len(ranked) >= 10:
                return ranked
    return ranked


def set_run_font(run, size: int, bold: bool = False, color: RGBColor | None = None, italic: bool = False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def add_horizontal_rule(paragraph) -> None:
    p = paragraph._p
    p_pr = p.get_or_add_pPr()
    p_bdr = p_pr.makeelement(
        qn("w:pBdr"),
        {},
    )
    bottom = p_bdr.makeelement(
        qn("w:bottom"),
        {
            qn("w:val"): "single",
            qn("w:sz"): "12",
            qn("w:space"): "1",
            qn("w:color"): "1F6F6A",
        },
    )
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def format_when(dt: datetime | None) -> str:
    if not dt:
        return "Date unavailable"
    return dt.strftime("%Y-%m-%d")


def write_document(sections: dict[str, list[dict]], generated_at: datetime) -> Path:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = header.add_run(generated_at.strftime("Daily News  |  %Y-%m-%d"))
        set_run_font(run, 9, color=MUTED)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Auto-generated briefing  |  sources: Google News, BBC, The Verge")
        set_run_font(run, 8, color=MUTED)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    style.paragraph_format.space_after = Pt(6)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("Daily News Briefing")
    set_run_font(run, 28, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(4)
    run = subtitle.add_run(generated_at.strftime("%A, %B %-d, %Y"))
    set_run_font(run, 14, color=TEAL)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(16)
    run = intro.add_run(
        generated_at.strftime("Generated automatically on %Y-%m-%d. ")
        + "Top 10 worldwide, China, and AI stories for the day."
    )
    set_run_font(run, 10, color=SLATE, italic=True)
    add_horizontal_rule(intro)

    for spec in SECTIONS:
        stories = sections[spec["key"]]
        heading = doc.add_paragraph()
        heading.paragraph_format.space_before = Pt(16)
        heading.paragraph_format.space_after = Pt(8)
        run = heading.add_run(spec["title"])
        set_run_font(run, 18, bold=True, color=NAVY)

        if not stories:
            empty = doc.add_paragraph()
            run = empty.add_run("No stories could be fetched for this section today.")
            set_run_font(run, 11, italic=True, color=MUTED)
            continue

        for index, story in enumerate(stories, start=1):
            item = doc.add_paragraph()
            item.paragraph_format.space_before = Pt(8)
            item.paragraph_format.space_after = Pt(2)
            item.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            run = item.add_run(f"{index}.  {story['title']}")
            set_run_font(run, 12, bold=True, color=NAVY)
            source = story["source"] or "Unknown source"
            run = item.add_run(f"  -  {source}")
            set_run_font(run, 11, color=TEAL)
            if story.get("title_zh"):
                title_zh = doc.add_paragraph()
                title_zh.paragraph_format.space_after = Pt(2)
                run = title_zh.add_run(story["title_zh"])
                set_run_font(run, 11, color=TEAL)

            meta = doc.add_paragraph()
            meta.paragraph_format.space_after = Pt(2)
            run = meta.add_run(f"Source: {source}  |  {format_when(story['published'])}")
            set_run_font(run, 9, color=MUTED)

            if story["summary"]:
                summary = doc.add_paragraph()
                summary.paragraph_format.space_after = Pt(2)
                run = summary.add_run(story["summary"])
                set_run_font(run, 11, color=SLATE)
            if story.get("summary_zh"):
                summary_zh = doc.add_paragraph()
                summary_zh.paragraph_format.space_after = Pt(2)
                run = summary_zh.add_run(story["summary_zh"])
                set_run_font(run, 11, italic=True, color=SLATE)

            if story["link"]:
                link_p = doc.add_paragraph()
                link_p.paragraph_format.space_after = Pt(8)
                run = link_p.add_run(story["link"])
                set_run_font(run, 9, color=LINK_BLUE)
                run.font.underline = True
                if run._element.rPr is not None:
                    r_id = doc.part.relate_to(
                        story["link"],
                        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                        is_external=True,
                    )
                    hyperlink = run._element.makeelement(
                        qn("w:hyperlink"),
                        {qn("r:id"): r_id},
                    )
                    run._element.addprevious(hyperlink)
                    hyperlink.append(run._element)

    output_path = OUTPUT_DIR / "Daily News.docx"
    for old in OUTPUT_DIR.glob("Daily News*.docx"):
        if old != output_path:
            old.unlink(missing_ok=True)
    doc.save(output_path)
    return output_path


def esc(value: str) -> str:
    return html.escape(value or "", quote=True)


def write_html(sections: dict[str, list[dict]], generated_at: datetime) -> Path:
    from web_page import render_site
    return render_site(sections, generated_at, OUTPUT_DIR, SECTIONS, format_when)


def main() -> int:
    generated_at = datetime.now(LOCAL_TZ)
    print(f"Generating daily news for {generated_at:%Y-%m-%d %H:%M %Z}")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    collected: dict[str, list[dict]] = {}
    for spec in SECTIONS:
        print(f"\n{spec['title']}")
        collected[spec["key"]] = collect_top10(spec["feeds"])
        print(f"  got {len(collected[spec['key']])} stories")

    print("\nTranslating English stories to Chinese")
    add_translations(collected)
    path = write_document(collected, generated_at)
    html_path = write_html(collected, generated_at)
    print(f"\nWrote {path}")
    print(f"Wrote {html_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
